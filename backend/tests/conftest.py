"""Test fixtures and mock document generators for pytest."""

import io
import pytest
import fitz  # PyMuPDF
import docx
from PIL import Image
from fastapi.testclient import TestClient
from app.main import app


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def sample_pdf_bytes():
    """Generates a valid, readable PDF document containing resume text."""
    doc = fitz.open()
    page = doc.new_page()
    text = (
        "John Doe\n"
        "Senior Software Engineer\n\n"
        "Summary:\n"
        "Experienced full-stack engineer proficient in Python, React, Node.js, and PostgreSQL.\n\n"
        "Experience:\n"
        "Software Engineer at TechCorp (2020-2023)\n"
        "- Built microservices using FastAPI, Docker, and Kubernetes.\n"
        "- Developed web applications with TypeScript, Next.js, and Tailwind CSS.\n\n"
        "Skills:\n"
        "Python, JavaScript, TypeScript, React, Next.js, Node.js, FastAPI, PostgreSQL, MongoDB, Docker, Git"
    )
    page.insert_text((50, 72), text, fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Generates a valid DOCX document containing paragraphs and tables."""
    doc = docx.Document()
    doc.add_heading("Jane Smith - Software Developer", 0)
    doc.add_paragraph("Summary: Passionate backend developer specialized in Java, Spring Boot, and AWS.")
    
    # Add skills table
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Category"
    hdr_cells[1].text = "Skills"

    row_cells = table.add_row().cells
    row_cells[0].text = "Languages"
    row_cells[1].text = "Java, Python, SQL, C++"

    row_cells2 = table.add_row().cells
    row_cells2[0].text = "Databases & Cloud"
    row_cells2[1].text = "PostgreSQL, Redis, AWS, Docker"

    doc_io = io.BytesIO()
    doc.save(doc_io)
    return doc_io.getvalue()


@pytest.fixture
def sample_png_bytes():
    """Generates a valid PNG image byte stream."""
    img = Image.new("RGB", (300, 100), color=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()
